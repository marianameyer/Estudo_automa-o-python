# Aplicativo para atualizar planilhas

"""
1- Entrar na planilha e extrair o CPF;
2- Entrar no site (https://consultcpf-devaprender.netlify.app/) e usar o CPF para pesquisar o 
status pagto;
3- Verificar se está "em dia" ou "atrasado";
4- Se "em dia", pegar data e método do pagto;
5 - Se "atrasado", status "pendente";
6- Inserir infos em nova planilha (nome, valor, cpf, vencimento, status, data pagto, método pagto);
7- Repetir até o último cliente
"""

# Importando as bibliotecas
import openpyxl
from selenium import webdriver  # Abre o navegador
from selenium.webdriver.common.by import By  # Encontrar e interagir com elementos da página
from time import sleep  # Pausas para carregar o site

# 1- Entrar na planilha e extrair o CPF;
planilha_clientes = openpyxl.load_workbook('/home/marianameyerf/app_atualizacao_planilhas/dados_clientes.xlsx')  # Abre a planilha
pagina_clientes = planilha_clientes['Sheet1']  # Acessa a página desejada

# Acessando o site
driver = webdriver.Chrome()
driver.get('https://consultcpf-devaprender.netlify.app/')

# Acessando as linhas
for linha in pagina_clientes.iter_rows(min_row=2, values_only=True):  # Dados começam da 2 linha
    nome, valor, cpf, vencimento = linha
    
# 2- Entrar no site (https://consultcpf-devaprender.netlify.app/) e usar o CPF para pesquisar o 
# status pagto;
    sleep(5)
    campo_pesquisa = driver.find_element(By.XPATH, "//input[@id='cpfInput']")
    sleep(1)
    campo_pesquisa.clear()  # Limpa o input do site
    campo_pesquisa.send_keys(cpf)
    sleep(1)

# 3- Verificar se está "em dia" ou "atrasado";
    botao_pesquisar = driver.find_element(By.XPATH, "//button[@class='btn btn-custom btn-lg btn-block mt-3']")
    sleep(1)
    botao_pesquisar.click()
    sleep(4)
    
# 
    status = driver.find_element(By.XPATH, "//span[@id='statusLabel']")

    # Verificação do status e coleta de infos
    if status.text == 'em dia':
        # 4- Se "em dia", pegar data e método do pagto;
        planilha_fechamento = openpyxl.load_workbook('/home/marianameyerf/app_atualizacao_planilhas/planilha_fechamento.xlsx')
        pagina_fechamento = planilha_fechamento['Sheet1']  # Acessa a planilha desejada

        # Obtendo infos
        data_pagamento = driver.find_element(By.XPATH, "//p[@id='paymentDate']")
        metodo_pagamento = driver.find_element(By.XPATH, "//p[@id='paymentMethod']")

        # Obtendo apenas a info necessária dos campos data e método pagto
        data_pgto_f = data_pagamento.text.split()[-1]
        metodo_pgto_f = metodo_pagamento.text.split()[-1]

        # Inserindo nova linha
        pagina_fechamento.append([nome, valor, cpf, vencimento, 'EM DIA', data_pgto_f, metodo_pgto_f])

        # Salvando as alterações na nova planilha
        planilha_fechamento.save('/home/marianameyerf/app_atualizacao_planilhas/planilha_fechamento.xlsx')

    else:
        # 5- Se "atrasado", status "pendente";
        planilha_fechamento = openpyxl.load_workbook('/home/marianameyerf/app_atualizacao_planilhas/planilha_fechamento.xlsx')
        pagina_fechamento = planilha_fechamento['Sheet1']  # Acessa a planilha desejada

        # Inserindo nova linha
        pagina_fechamento.append([nome, valor, cpf, vencimento, 'PENDENTE'])

        # Salvando as alterações na nova planilha
        planilha_fechamento.save('/home/marianameyerf/app_atualizacao_planilhas/planilha_fechamento.xlsx')
        
# 6- Inserir infos em nova planilha (nome, valor, cpf, vencimento, status, data pagto, método pagto);
# 7- Repetir até o último cliente